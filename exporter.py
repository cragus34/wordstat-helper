"""
Генерация красивого Excel-файла из результатов запроса.
Включает заголовок с исходной фразой (маской) и метаданными.
"""
import io
import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def sanitize_sheet_title(title: str) -> str:
    """Очищает строку для использования в названии листа Excel (макс 31 символ, без :\\/?*[])."""
    clean = re.sub(r'[:\\/?*\[\]]', '', title).strip()
    return clean[:30] if clean else "Семантика"


def build_xlsx(seed_phrase: str, results: list[dict], created_at: str = "") -> bytes:
    """
    Принимает seed_phrase и список {"phrase", "base_freq", "exact_freq"}.
    Возвращает байты .xlsx файла со стильным оформлением.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sanitize_sheet_title(seed_phrase)
    ws.views.sheetView[0].showGridLines = True

    # Стили
    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    meta_font = Font(name="Calibri", size=10, italic=True, color="6B7280")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    
    thin_border_side = Side(style="thin", color="E5E7EB")
    border_data = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    # ── Блок 1: Заголовок отчёта ──────────────────────────────────────────
    ws.cell(row=1, column=1, value=f"Маска запроса: «{seed_phrase}»").font = title_font
    
    date_str = created_at or datetime.now().strftime("%Y-%m-%d %H:%M")
    ws.cell(
        row=2, column=1, 
        value=f"Дата сбора: {date_str}  |  Всего фраз: {len(results)}  |  Сортировка: по точной частотности [!фраза]"
    ).font = meta_font

    # ── Блок 2: Заголовки таблицы ─────────────────────────────────────────
    table_headers = ["Ключевая фраза", "Базовая частотность", "Точная частотность [!фраза]"]
    start_row = 4

    for col_idx, header_text in enumerate(table_headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[start_row].height = 26

    # ── Блок 3: Строки данных ─────────────────────────────────────────────
    data_font = Font(name="Calibri", size=11)
    zebra_fill = PatternFill("solid", fgColor="F8FAFC")

    for i, item in enumerate(results):
        r_idx = start_row + 1 + i
        exact_value = item.get("exact_freq")
        if exact_value is None:
            exact_value = item.get("super_exact_freq")

        c1 = ws.cell(row=r_idx, column=1, value=item["phrase"])
        c2 = ws.cell(row=r_idx, column=2, value=item["base_freq"])
        c3 = ws.cell(row=r_idx, column=3, value=exact_value)

        c1.font = data_font
        c2.font = data_font
        c3.font = Font(name="Calibri", size=11, bold=True, color="1D4ED8")

        c1.alignment = Alignment(horizontal="left", vertical="center")
        c2.alignment = Alignment(horizontal="right", vertical="center")
        c3.alignment = Alignment(horizontal="right", vertical="center")

        # Чередование строк
        if i % 2 == 1:
            c1.fill = zebra_fill
            c2.fill = zebra_fill
            c3.fill = zebra_fill

        c1.border = border_data
        c2.border = border_data
        c3.border = border_data

    # Ширина колонок
    max_phrase_len = max([len(item["phrase"]) for item in results], default=25)
    ws.column_dimensions["A"].width = max(max_phrase_len + 4, 35)
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
